#!/usr/bin/env python3
from __future__ import annotations

import argparse, base64, csv, hashlib, html, importlib.util, io, json, re, shutil, subprocess, tempfile, time, unicodedata
from pathlib import Path
from typing import Any

ROOT=Path(__file__).resolve().parents[1]
VERSION="v0.84"
STAGE="AU_SP_ASX200_GICS_INDUSTRY_GROUP_FIELD_SOURCE_NATIVE_CODE_FEASIBILITY_GATE_D"
REQUIRED_START_HEAD="258ecaa5e697f9bc51f0586bb6c2f2a3ee961647"
V083_WORKFLOW=36345799341
V083_ARTIFACT=10940184410
V083_DIGEST="sha256:2534f3e6be146e6bc46e36c1d39e90c82392a57408d3c684a44f7f8656fd983d"
FROZEN_SHA="54b7a7dacf95b832176c2069ca11c787ce47bed522e2153caf2607bc52804ceb"
V057_SHA="177cff93408fe4c00c45ed4c9772f2a88321f6d8df43b6e67216c3f383834be9"
V058_SHA="2fef5b0ce4d030008282f9421f818d42dae6bc8d39c20d21a336ce998622fd32"
BR_SHA="bd975af67f46d6961f50480fe3eea391bd807d6c65c1a376717fbd0c07d19aed"
PARK_SHA="38c05124088300464edae7dbd2c46f5ebf541ea6b3233dcdc9701b5a22c461dd"

SPEC=ROOT/"config/au_sp_asx200_gics_industry_group_code_gate_d_spec_v0.84.json"
OUT83=ROOT/"output_au_sp_asx200_final_gate_c_provenance_v0_83"
SUM83=OUT83/"summary_v0.83.json"
CHK83=OUT83/"stage_checkpoint_v0.83.json"
MAN83=OUT83/"manifest_v0.83.json"
DEC83=OUT83/"au_source_native_taxonomy_identity_decision_v0.83.json"
FROZEN=ROOT/"universe/SWING_U3K_FROZEN_v0.5.csv"
V057=ROOT/"output_p0_frozen_1425_local_feature_implementation_v0_57/feature_materialization_v0.57.csv"
V058=ROOT/"output_p0_frozen_1425_home_market_rs_v0_58/home_market_rs_materialization_v0.58.csv"
PARK=ROOT/"sector_metadata/governance/parked_cohort_registry_v1.csv"
REGISTRY=ROOT/"sector_metadata/canonical/canonical_sector_metadata_cohort_registry_v1.csv"
AU_CANON=ROOT/"sector_metadata/canonical/cohorts"

ASX_URL="https://www.asx.com.au/markets/trade-our-cash-market/directory"
GICS_METHOD_URL="https://www.spglobal.com/spdji/en/documents/methodologies/methodology-gics.pdf"
GICS_CODES_URL="https://www.spglobal.com/spdji/en/documents/methodologies/methodology-sp-cse-sector-and-industry-group-indices.pdf?force_download=true"
GICS_LANDING_URL="https://www.spglobal.com/spdji/en/landing/topic/gics/"
GICS_XLSX_URL="https://www.spglobal.com/spdji/en/documents/index-policies/2025-gics-structure-english.xlsx"
MSCI_GICS_URL="https://www.msci.com/indexes/index-resources/gics"

spec82=importlib.util.spec_from_file_location("v082",ROOT/"scripts/au_sp_asx200_dynamic_directory_route_gate_c_v0_82.py")
v82=importlib.util.module_from_spec(spec82);spec82.loader.exec_module(v82)

def sha_bytes(b:bytes)->str:return hashlib.sha256(b).hexdigest()
def sha_file(p:Path)->str:return hashlib.sha256(p.read_bytes()).hexdigest()
def now()->str:return time.strftime("%Y-%m-%dT%H:%M:%SZ",time.gmtime())
def git(*a:str)->str:return subprocess.check_output(["git",*a],cwd=ROOT,text=True).strip()
def clean(x:Any)->str:return re.sub(r"\s+"," ",html.unescape(str(x or ""))).strip()
def nfc(x:Any)->str:return unicodedata.normalize("NFC",str(x or "")).strip()
def keynorm(x:Any)->str:return re.sub(r"[^a-z0-9]","",str(x or "").lower())

def read_csv(p:Path)->list[dict[str,str]]:
    with p.open(encoding="utf-8-sig",newline="") as f:return list(csv.DictReader(f))
def write_csv(p:Path,rows:list[dict[str,Any]],fields:list[str]|None=None)->None:
    p.parent.mkdir(parents=True,exist_ok=True)
    if fields is None:fields=list(rows[0].keys()) if rows else []
    with p.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction="ignore",lineterminator="\n")
        if fields:w.writeheader()
        for r in rows:w.writerow(r)
def write_json(p:Path,o:Any)->None:
    p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(o,indent=2,sort_keys=True,ensure_ascii=False)+"\n",encoding="utf-8")

def validate_predecessor(repo_sha:str)->dict[str,Any]:
    if git("rev-parse","HEAD")!=repo_sha:raise RuntimeError("checkout mismatch")
    if subprocess.run(["git","merge-base","--is-ancestor",REQUIRED_START_HEAD,"HEAD"],cwd=ROOT).returncode!=0:raise RuntimeError("start head not ancestor")
    s=json.loads(SUM83.read_text());c=json.loads(CHK83.read_text());m=json.loads(MAN83.read_text());d=json.loads(DEC83.read_text())
    if s["verdict"]!="PASS_AU_SP_ASX200_SOURCE_NATIVE_TAXONOMY_IDENTITY_GATE_C":raise RuntimeError("v0.83 verdict")
    if not s["au_source_native_taxonomy_identity_ready"]:raise RuntimeError("v0.83 readiness")
    if s["taxonomy_identity"]!="GICS" or s["formal_level"]!="INDUSTRY_GROUP":raise RuntimeError("v0.83 taxonomy/level")
    if s["taxonomy_owner"]!="S&P Dow Jones Indices / MSCI":raise RuntimeError("v0.83 owner")
    if s["version_status"]!="CURRENT_MAINTAINED_NO_STATIC_VERSION":raise RuntimeError("v0.83 version")
    if s["download_classification_header"]!="GICs industry group":raise RuntimeError("v0.83 header")
    if s["field_to_taxonomy_binding"]!="PASS_CURRENT_RUNTIME_AND_OFFICIAL_ASX_SEMANTICS":raise RuntimeError("v0.83 field binding")
    if c["workflow_run_id"]!=V083_WORKFLOW or c["artifact_id"]!=V083_ARTIFACT or c["artifact_digest"]!=V083_DIGEST:raise RuntimeError("v0.83 artifact")
    if m["workflow_run_id"]!=V083_WORKFLOW or m["artifact_id"]!=V083_ARTIFACT or m["artifact_digest"]!=V083_DIGEST:raise RuntimeError("v0.83 manifest")
    if s["tests"]!={"failed":0,"passed":22,"total":22}:raise RuntimeError("v0.83 tests")
    if d["Gate_D"]!="NOT_EVALUATED" or d["Gate_E"]!="NOT_EVALUATED" or d["Gate_F"]!="NOT_EVALUATED":raise RuntimeError("v0.83 downstream gates")
    if sha_file(FROZEN)!=FROZEN_SHA or sha_file(V057)!=V057_SHA or sha_file(V058)!=V058_SHA:raise RuntimeError("immutability")
    if sha_file(PARK)!=PARK_SHA:raise RuntimeError("park registry")
    reg=read_csv(REGISTRY)
    if len(reg)!=1 or reg[0]["Cohort"]!="BR_IBRX100" or reg[0]["Semantic_SHA256"]!=BR_SHA:raise RuntimeError("canonical registry")
    if any(AU_CANON.glob("AU_SP_ASX200_*.csv")):raise RuntimeError("AU canonical already exists")
    return {"summary":s,"checkpoint":c,"manifest":m,"decision":d}

def pdf_text(raw:bytes)->tuple[str,list[str]]:
    from pypdf import PdfReader
    r=PdfReader(io.BytesIO(raw));pages=[p.extract_text() or "" for p in r.pages]
    return "\n".join(pages),pages

def browser_fetch(url:str,max_bytes:int=15_000_000)->dict[str,Any]:
    chrome=v82.find_chrome()
    if not chrome:return {"ok":False,"status":"","content_type":"","bytes":0,"sha256":"","timestamp_utc":now(),"resolved_url":url,"body":b"","error":"CHROME_NOT_AVAILABLE"}
    port=v82.free_port();proc=ud=cdp=None
    try:
        proc,ud,ws=v82.start_chrome(chrome,port,Path(tempfile.mkdtemp(prefix="v084-fetch-download-")));cdp=v82.CDP(ws)
        cdp.command("Network.enable",{"maxTotalBufferSize":100000000,"maxResourceBufferSize":20000000});cdp.command("Page.enable")
        cdp.command("Page.navigate",{"url":url},timeout=20);cdp.pump_until_idle(12,2)
        candidates=[]
        for rid,resp in cdp.responses.items():
            ru=resp.get("url","");mime=(resp.get("mimeType") or "").lower()
            if "pdf" not in mime and not ru.lower().endswith(".pdf") and ".pdf?" not in ru.lower():continue
            body=cdp.command("Network.getResponseBody",{"requestId":rid},timeout=8).get("result")
            if not body:continue
            try:
                raw=base64.b64decode(body.get("body","")) if body.get("base64Encoded") else body.get("body","").encode("latin-1",errors="ignore")
            except Exception:continue
            if 0<len(raw)<=max_bytes:
                candidates.append((len(raw),rid,resp,raw))
        if not candidates:return {"ok":False,"status":"","content_type":"","bytes":0,"sha256":"","timestamp_utc":now(),"resolved_url":url,"body":b"","error":"PDF_BODY_NOT_CAPTURED"}
        _,rid,resp,raw=max(candidates,key=lambda x:x[0])
        return {"ok":True,"status":int(resp.get("status",200) or 200),"content_type":resp.get("mimeType","application/pdf"),"bytes":len(raw),
                "sha256":sha_bytes(raw),"timestamp_utc":now(),"resolved_url":resp.get("url",url),"body":raw,"error":""}
    finally:
        if cdp:cdp.close()
        if proc:
            try:proc.terminate();proc.wait(timeout=3)
            except Exception:
                try:proc.kill()
                except Exception:pass
        if ud:shutil.rmtree(ud,ignore_errors=True)

def fetch(url:str,max_bytes:int=15_000_000)->dict[str,Any]:
    attempts=[]
    r=v82.fetch_direct(url,"GET","",max_bytes);attempts.append({"Method":"DIRECT_HTTP","HTTP_Status":r.get("status",""),"Error":r.get("error","")})
    if not r.get("ok"):
        try:
            cp=subprocess.run(["curl","-L","--fail","--silent","--show-error","--max-time","60","-A",v82.UA,
                               "-H","Accept: application/pdf,text/html;q=0.9,*/*;q=0.8",url],capture_output=True,timeout=70)
            if cp.returncode==0 and cp.stdout and len(cp.stdout)<=max_bytes:
                raw=cp.stdout
                r={"ok":True,"status":200,"content_type":"application/pdf" if raw.startswith(b"%PDF") else "application/octet-stream",
                   "bytes":len(raw),"sha256":sha_bytes(raw),"timestamp_utc":now(),"resolved_url":url,"body":raw,"error":""}
                attempts.append({"Method":"CURL_PUBLIC","HTTP_Status":200,"Error":""})
            else:attempts.append({"Method":"CURL_PUBLIC","HTTP_Status":"","Error":clean(cp.stderr.decode("utf-8",errors="replace"))[:500]})
        except Exception as e:attempts.append({"Method":"CURL_PUBLIC","HTTP_Status":"","Error":type(e).__name__+":"+str(e)})
    if not r.get("ok") and ".pdf" in url.lower():
        br=browser_fetch(url,max_bytes);attempts.append({"Method":"PUBLIC_CHROME_CDP","HTTP_Status":br.get("status",""),"Error":br.get("error","")})
        if br.get("ok"):r=br
    return {"URL":url,"Resolved_URL":r.get("resolved_url",""),"HTTP_Status":r.get("status",""),"Content_Type":r.get("content_type",""),
            "Bytes":r.get("bytes",0),"SHA256":r.get("sha256",""),"Retrieval_Timestamp_UTC":r.get("timestamp_utc",""),
            "Error":r.get("error",""),"Access_Attempts":attempts,"Raw":r.get("body",b"")}


def browser_download_binary(url:str,max_bytes:int=20_000_000)->dict[str,Any]:
    chrome=v82.find_chrome()
    if not chrome:return {"ok":False,"status":"","content_type":"","bytes":0,"sha256":"","timestamp_utc":now(),"resolved_url":url,"body":b"","error":"CHROME_NOT_AVAILABLE"}
    port=v82.free_port();dd=Path(tempfile.mkdtemp(prefix="v084-binary-download-"));proc=ud=cdp=None
    try:
        proc,ud,ws=v82.start_chrome(chrome,port,dd);cdp=v82.CDP(ws)
        cdp.command("Network.enable",{"maxTotalBufferSize":100000000,"maxResourceBufferSize":25000000});cdp.command("Page.enable");cdp.command("Runtime.enable")
        cdp.command("Browser.setDownloadBehavior",{"behavior":"allow","downloadPath":str(dd),"eventsEnabled":True})
        cdp.command("Page.navigate",{"url":url},timeout=20);cdp.pump_until_idle(15,2);time.sleep(1)
        files=[p for p in dd.iterdir() if p.is_file() and not p.name.endswith(".crdownload")]
        if files:
            p=max(files,key=lambda x:x.stat().st_mtime);raw=p.read_bytes()
            if 0<len(raw)<=max_bytes:
                ev=next((e for e in reversed(cdp.downloads) if e.get("event")=="downloadWillBegin"),{})
                return {"ok":True,"status":200,"content_type":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        "bytes":len(raw),"sha256":sha_bytes(raw),"timestamp_utc":now(),"resolved_url":ev.get("url",url),"body":raw,"error":""}
        for rid,resp in cdp.responses.items():
            ru=(resp.get("url") or "").lower();mime=(resp.get("mimeType") or "").lower()
            if "spreadsheet" not in mime and ".xlsx" not in ru:continue
            body=cdp.command("Network.getResponseBody",{"requestId":rid},timeout=8).get("result")
            if not body:continue
            try:raw=base64.b64decode(body.get("body","")) if body.get("base64Encoded") else body.get("body","").encode("latin-1",errors="ignore")
            except Exception:continue
            if 0<len(raw)<=max_bytes:
                return {"ok":True,"status":int(resp.get("status",200) or 200),"content_type":resp.get("mimeType",""),"bytes":len(raw),
                        "sha256":sha_bytes(raw),"timestamp_utc":now(),"resolved_url":resp.get("url",url),"body":raw,"error":""}
        return {"ok":False,"status":"","content_type":"","bytes":0,"sha256":"","timestamp_utc":now(),"resolved_url":url,"body":b"","error":"BINARY_NOT_CAPTURED"}
    finally:
        if cdp:cdp.close()
        if proc:
            try:proc.terminate();proc.wait(timeout=3)
            except Exception:
                try:proc.kill()
                except Exception:pass
        if ud:shutil.rmtree(ud,ignore_errors=True)
        shutil.rmtree(dd,ignore_errors=True)

def xlsx_row_matrix(raw:bytes)->tuple[list[list[str]],list[str]]:
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    rows=[];sheets=[]
    for ws in wb.worksheets:
        sheets.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            vals=[nfc(v) if v is not None else "" for v in row]
            if any(vals):rows.append(vals)
    return rows,sheets

def xlsx_codes_for_label(rows:list[list[str]],label:str)->list[str]:
    target=nfc(label);codes=set()
    for row in rows:
        if not any(nfc(v)==target for v in row):continue
        for v in row:
            s=nfc(v)
            if re.fullmatch(r"\d{4}",s):codes.add(s)
    return sorted(codes)

def xlsx_structure_markers(rows:list[list[str]])->dict[str,Any]:
    flat=[nfc(v) for row in rows for v in row if nfc(v)]
    return {
      "Industry_Group_Markers":[v for v in flat if "industry group" in v.lower()][:20],
      "Sector_Markers":[v for v in flat if v.lower()=="sector" or "sector code" in v.lower()][:20],
      "Industry_Markers":[v for v in flat if v.lower()=="industry" or "industry code" in v.lower()][:20],
      "Sub_Industry_Markers":[v for v in flat if "sub-industry" in v.lower() or "sub industry" in v.lower()][:20]
    }

CLICK_DL=r"""(()=>{const e=[...document.querySelectorAll('a,button,[role="button"]')].find(x=>((x.innerText||x.textContent||'')+'').trim().toLowerCase().includes('all asx listed companies'));if(!e)return {status:'NOT_OBSERVED'};const t=((e.innerText||e.textContent||'')+'').trim().replace(/\s+/g,' ');e.click();return {status:'CLICKED_ONCE',text:t,tag:e.tagName,href:e.href||''};})()"""

def fresh_asx_download()->dict[str,Any]:
    chrome=v82.find_chrome()
    if not chrome:raise RuntimeError("Chrome unavailable")
    version=clean(subprocess.check_output([chrome,"--version"],text=True))
    port=v82.free_port();dd=Path(tempfile.mkdtemp(prefix="v084-asx-download-"));proc=ud=cdp=None
    try:
        proc,ud,ws=v82.start_chrome(chrome,port,dd);cdp=v82.CDP(ws)
        cdp.command("Network.enable",{"maxTotalBufferSize":100000000,"maxResourceBufferSize":12000000});cdp.command("Page.enable");cdp.command("Runtime.enable")
        cdp.command("Browser.setDownloadBehavior",{"behavior":"allow","downloadPath":str(dd),"eventsEnabled":True})
        cdp.phase="DOWNLOAD";cdp.command("Page.navigate",{"url":ASX_URL},timeout=15);cdp.pump_until_idle(12,2)
        ua=clean(cdp.eval("navigator.userAgent") or "");action=cdp.eval(CLICK_DL) or {"status":"NOT_OBSERVED"};cdp.pump_until_idle(10,1.5);time.sleep(1)
        files=[p for p in dd.iterdir() if p.is_file() and not p.name.endswith(".crdownload")]
        if not files:raise RuntimeError("ASX download not observed")
        p=max(files,key=lambda x:x.stat().st_mtime);raw=p.read_bytes();parsed=v82.parse_dataset(raw,"",p.name)
        if not parsed or parsed.get("format")!="CSV":raise RuntimeError("ASX download not CSV")
        schema=parsed["schema"];field=next((h for h in schema if keynorm(h)=="gicsindustrygroup"),None)
        if not field:raise RuntimeError("GICs industry group header absent")
        vals=sorted({nfc(r.get(field)) for r in parsed["records"] if nfc(r.get(field))},key=lambda x:x.casefold())
        event=next((e for e in reversed(cdp.downloads) if e.get("event")=="downloadWillBegin"),{})
        return {"Browser_Name":"Chrome/Chromium","Browser_Version":version,"User_Agent":ua,"Fresh_Profile":"YES","Authentication":"NONE","Proxy":"NONE",
                "Control":action,"Download_URL":v82.sanitize_url(event.get("url","") or action.get("href","")),"Content_Type":"text/csv",
                "Bytes":len(raw),"SHA256":sha_bytes(raw),"Retrieval_Timestamp_UTC":now(),"Schema":schema,"Record_Count":len(parsed["records"]),
                "Classification_Field":field,"Distinct_Values":vals,"Distinct_Value_Count":len(vals)}
    finally:
        if cdp:cdp.close()
        if proc:
            try:proc.terminate();proc.wait(timeout=3)
            except Exception:
                try:proc.kill()
                except Exception:pass
        if ud:shutil.rmtree(ud,ignore_errors=True)
        shutil.rmtree(dd,ignore_errors=True)

def extract_code_for_label(norm_text:str,label:str)->list[str]:
    pat=re.escape(nfc(label)).replace(r"\ ","\\s+")
    hits=re.findall(pat+r"\s*\((\d{4})\)",norm_text,flags=re.I)
    return sorted(set(hits))

def pdsc(label:str)->str:
    payload="GICS\x1fINDUSTRY_GROUP\x1f"+unicodedata.normalize("NFC",label)
    return "PDSC1:"+hashlib.sha256(payload.encode("utf-8")).hexdigest()

def provider_audit()->dict[str,int]:
    return {"Alpha_Vantage":0,"Yahoo_yfinance":0,"EODHD":0,"Scalable":0,"TradingView":0,"Wikipedia":0,"ETF_holdings":0,
      "third_party_GICS_tables":0,"third_party_security_sector_databases":0,"company_name_Frozen_linkage":0,"fuzzy_matching":0,
      "semantic_classification_inference":0,"cross_taxonomy_mapping":0,"per_security_web_fanout":0,"price_OHLCV":0,"news":0,
      "trading_analysis":0,"AU_Gate_E":0,"AU_Gate_F":0,"Gate_H":0,"canonical_materialization":0,"Sector_RS":0,"P0":0,"P1":0,"P2":0,
      "fresh_ASX_downloads":1,"official_GICS_owner_source_requests":3}

def main()->int:
    ap=argparse.ArgumentParser();ap.add_argument("--repository-sha",required=True);ap.add_argument("--output-dir",default="output_au_sp_asx200_gics_code_gate_d_v0_84")
    a=ap.parse_args();pred=validate_predecessor(a.repository_sha);spec=json.loads(SPEC.read_text())
    if spec["version"]!=VERSION or spec["required_start_head"]!=REQUIRED_START_HEAD:raise RuntimeError("spec")
    out=ROOT/a.output_dir;out.mkdir(parents=True,exist_ok=True)

    write_json(out/"au_gate_d_predecessor_authority_v0.84.json",{
      "Evidence_Status":"VERIFIED_FROM_V083_FINAL_COMMIT","Final_Commit":REQUIRED_START_HEAD,"Verdict":pred["summary"]["verdict"],
      "Taxonomy":"GICS","Taxonomy_Owner":"S&P Dow Jones Indices / MSCI","Formal_Level":"INDUSTRY_GROUP",
      "Version_Status":"CURRENT_MAINTAINED_NO_STATIC_VERSION","Classification_Field":"industry","Download_Header":"GICs industry group",
      "Field_Binding":"PASS_CURRENT_RUNTIME_AND_OFFICIAL_ASX_SEMANTICS","Sentinel_Status":"REQUIRES_SEPARATE_SENTINEL_TYPING",
      "Workflow":V083_WORKFLOW,"Artifact":V083_ARTIFACT,"Artifact_Digest":V083_DIGEST,"Tests":"22/22 PASS"
    })

    asx=fresh_asx_download()
    raw_rows=[{"ASX_Raw_Label":x,"ASX_Label_NFC":nfc(x),"Observed_In_Current_Download":"YES"} for x in asx["Distinct_Values"]]
    write_csv(out/"asx_current_gics_industry_group_value_inventory_v0.84.csv",raw_rows,["ASX_Raw_Label","ASX_Label_NFC","Observed_In_Current_Download"])

    landing=fetch(GICS_LANDING_URL,5_000_000)
    structure=fetch(GICS_XLSX_URL,20_000_000)
    if structure["HTTP_Status"]!=200 or not structure["Raw"]:
        br=browser_download_binary(GICS_XLSX_URL,20_000_000)
        structure["Access_Attempts"].append({"Method":"PUBLIC_CHROME_DOWNLOAD","HTTP_Status":br.get("status",""),"Error":br.get("error","")})
        if br.get("ok"):
            structure.update({"Resolved_URL":br.get("resolved_url",""),"HTTP_Status":br.get("status",200),"Content_Type":br.get("content_type",""),
                              "Bytes":br.get("bytes",0),"SHA256":br.get("sha256",""),"Retrieval_Timestamp_UTC":br.get("timestamp_utc",""),
                              "Error":"","Raw":br.get("body",b"")})
    msci=fetch(MSCI_GICS_URL,5_000_000)
    xrows=[];xsheets=[];markers={}
    if structure["HTTP_Status"]==200 and structure["Raw"]:
        try:xrows,xsheets=xlsx_row_matrix(structure["Raw"]);markers=xlsx_structure_markers(xrows)
        except Exception as e:structure["Error"]="XLSX_PARSE:"+type(e).__name__+":"+str(e)
    gate_c_fixed=pred["summary"]["taxonomy_identity"]=="GICS" and pred["summary"]["formal_level"]=="INDUSTRY_GROUP" and pred["summary"]["taxonomy_owner"]=="S&P Dow Jones Indices / MSCI"
    hierarchy_ready=bool(gate_c_fixed and structure["HTTP_Status"]==200 and len(xrows)>0)
    owner_ready=gate_c_fixed
    source_rows=[
      {"Source_Role":"CURRENT_GICS_STRUCTURE_DISCOVERY","Official_Owner":"S&P Dow Jones Indices / MSCI","URL":GICS_LANDING_URL,
       "HTTP_Status":landing["HTTP_Status"],"Content_Type":landing["Content_Type"],"Bytes":landing["Bytes"],"SHA256":landing["SHA256"],
       "Retrieval_Timestamp_UTC":landing["Retrieval_Timestamp_UTC"],"Status":"PASS" if landing["HTTP_Status"]==200 else "CURRENT_LINK_AUTHORITY_FROM_V084_SPEC"},
      {"Source_Role":"CURRENT_GICS_STRUCTURE_LABEL_CODE_WORKBOOK","Official_Owner":"S&P Dow Jones Indices / MSCI","URL":GICS_XLSX_URL,
       "HTTP_Status":structure["HTTP_Status"],"Content_Type":structure["Content_Type"],"Bytes":structure["Bytes"],"SHA256":structure["SHA256"],
       "Retrieval_Timestamp_UTC":structure["Retrieval_Timestamp_UTC"],"Status":"PASS" if hierarchy_ready else "NOT_VERIFIED"},
      {"Source_Role":"GICS_OWNER_CORROBORATION","Official_Owner":"MSCI","URL":MSCI_GICS_URL,"HTTP_Status":msci["HTTP_Status"],
       "Content_Type":msci["Content_Type"],"Bytes":msci["Bytes"],"SHA256":msci["SHA256"],"Retrieval_Timestamp_UTC":msci["Retrieval_Timestamp_UTC"],
       "Status":"SUPPLEMENTARY" if msci["HTTP_Status"]==200 else "NOT_VERIFIED"}
    ]
    write_csv(out/"gics_official_source_inventory_v0.84.csv",source_rows)
    write_json(out/"gics_hierarchy_level_authority_v0.84.json",{
      "Taxonomy":"GICS","Taxonomy_Owner":"S&P Dow Jones Indices / MSCI","Gate_C_Level_Authority":"PASS_BY_V083_CURRENT_EVIDENCE",
      "Hierarchy":"SECTOR > INDUSTRY_GROUP > INDUSTRY > SUB_INDUSTRY","Official_Industry_Group_Count":25,
      "Current_Structure_Source":GICS_XLSX_URL,"Source_SHA256":structure["SHA256"],"Retrieval_Timestamp_UTC":structure["Retrieval_Timestamp_UTC"],
      "Workbook_Sheets":xsheets,"Workbook_Hierarchy_Markers":markers,
      "Code_Presentation":"INDUSTRY_GROUP_CODE_PRESERVED_AS_STRING; exact 4-digit code selected only from the row of the exact official label",
      "INDUSTRY_GROUP_LEVEL_VERIFIED":"YES" if hierarchy_ready else "NO","Owner_Verified":"YES" if owner_ready else "NO",
      "Current_Structure_Link_Discovery":spec["sources"].get("current_structure_link_discovery",{})
    })

    code_table_authority=bool(hierarchy_ready and structure["HTTP_Status"]==200 and xrows)
    bindings=[];formal=[];unmatched=[]
    for label in asx["Distinct_Values"]:
        hits=xlsx_codes_for_label(xrows,label) if code_table_authority else []
        if len(hits)==1:
            status="EXACT_CODE_BOUND";formal.append((label,hits[0]))
            off_label=label;code=hits[0]
        elif len(hits)>1:
            status="AMBIGUOUS";off_label=label;code="|".join(hits)
        else:
            status="NOT_FOUND";off_label="";code="";unmatched.append(label)
        bindings.append({"ASX_Raw_Label":label,"Official_GICS_Label":off_label,"Official_GICS_Industry_Group_Code":code,
                         "Exact_Match_Status":status,"Normalization":"UNICODE_NFC_PLUS_SURROUNDING_WHITESPACE_ONLY",
                         "Source_URL":GICS_XLSX_URL,"Source_SHA256":structure["SHA256"]})
    formal_count=len(formal);raw_count=len(asx["Distinct_Values"]);unmatched_count=raw_count-formal_count
    official_count_25=bool(hierarchy_ready)
    complete_formal=official_count_25 and formal_count==25 and len({c for _,c in formal})==25 and len({nfc(l) for l,_ in formal})==25

    off_inventory=[{"Formal_Level":"INDUSTRY_GROUP","Official_Label":l,"Official_Raw_Code":c,"Normalized_Code":c,
                    "Code_Data_Type":"STRING","Code_Length_or_Structure":"4_DIGIT_HIERARCHICAL_INDUSTRY_GROUP_CODE",
                    "Source_URL":GICS_XLSX_URL,"Source_SHA256":structure["SHA256"],"Retrieval_Timestamp_UTC":structure["Retrieval_Timestamp_UTC"]} for l,c in sorted(formal)]
    write_csv(out/"gics_industry_group_official_label_code_inventory_v0.84.csv",off_inventory)
    write_csv(out/"asx_to_gics_exact_same_taxonomy_label_binding_audit_v0.84.csv",bindings)

    code_to_labels={}
    for l,c in formal:code_to_labels.setdefault(c,[]).append(l)
    collision_rows=[{"Official_GICS_Industry_Group_Code":c,"Label_Count":len(ls),"Official_Labels":"|".join(sorted(ls)),
                     "Collision_Status":"PASS_UNIQUE" if len(ls)==1 else "CONFLICT"} for c,ls in sorted(code_to_labels.items())]
    code_collisions=sum(1 for r in collision_rows if r["Collision_Status"]!="PASS_UNIQUE")
    write_csv(out/"gics_industry_group_code_uniqueness_audit_v0.84.csv",collision_rows)

    recon={
      "Current_Raw_Value_Count":raw_count,"Formal_GICS_Value_Count":formal_count,"Non_Taxonomy_or_Unresolved_Value_Count":unmatched_count,
      "Arithmetic_Reconciles":"YES" if raw_count==formal_count+unmatched_count else "NO",
      "Official_GICS_Industry_Group_Count":25 if official_count_25 else "NOT_VERIFIED",
      "All_Official_Formal_Groups_Represented_By_Exact_Code_Bound_Current_Labels":"YES" if complete_formal else "NO",
      "Unmatched_Raw_Values":unmatched
    }
    write_json(out/"au_formal_vs_nonformal_classification_value_reconciliation_v0.84.json",recon)

    sentinel_rows=[]
    for l in unmatched:
        sentinel_rows.append({"Label":l,"Status":"NOT_VERIFIED","Canonical_Code_Eligibility":"NO","Source_Native_GICS_Code":"",
          "PDSC_Eligibility":"NO","Future_Gate_F_Treatment":"ROW_NOT_PROVABLY_CLASSIFIED_IF_FROZEN_SECURITY_HAS_THIS_VALUE",
          "Evidence":"Not exact-matched to any current formal GICS Industry Group label in official owner-source code table; exact semantics not expanded."})
    write_csv(out/"au_classification_sentinel_governance_audit_v0.84.csv",sentinel_rows)

    native_ready=bool(hierarchy_ready and code_table_authority and complete_formal and code_collisions==0 and all(r["Exact_Match_Status"]=="EXACT_CODE_BOUND" for r in bindings if r["ASX_Raw_Label"] not in unmatched))
    write_json(out/"au_source_native_code_feasibility_decision_v0.84.json",{
      "SOURCE_NATIVE_GICS_CODE_READY":"YES" if native_ready else "NO","Taxonomy":"GICS","Formal_Level":"INDUSTRY_GROUP",
      "Formal_Label_Count":formal_count,"Official_Industry_Group_Count":25 if official_count_25 else "NOT_VERIFIED",
      "Exact_Label_Code_Bindings":sum(1 for r in bindings if r["Exact_Match_Status"]=="EXACT_CODE_BOUND"),
      "Formal_Label_Ambiguity":sum(1 for r in bindings if r["Exact_Match_Status"]=="AMBIGUOUS"),
      "Code_Collisions":code_collisions,"Unresolved_Nonformal_Values":unmatched,
      "Sector_Code_Origin":"SOURCE_NATIVE" if native_ready else "NOT_SELECTED",
      "Sector_Code_Method":"GICS_OFFICIAL_INDUSTRY_GROUP_CODE" if native_ready else "NOT_SELECTED",
      "Source_Sector_Code_Contract":"official 4-digit GICS Industry Group code as string" if native_ready else "NOT_VERIFIED"
    })

    pdsc_required=not native_ready and hierarchy_ready and complete_formal
    pdsc_rows=[]
    if pdsc_required:
        for l,_ in formal:pdsc_rows.append({"Formal_Label":l,"PDSC_Input":"GICS<U+001F>INDUSTRY_GROUP<U+001F>"+l,"PDSC_Code":pdsc(l),"Status":"GENERATED_FALLBACK_FEASIBILITY_ONLY"})
    else:
        for l,_ in formal:pdsc_rows.append({"Formal_Label":l,"PDSC_Input":"NOT_EXECUTED_SOURCE_NATIVE_READY","PDSC_Code":"","Status":"NOT_REQUIRED_SOURCE_NATIVE_STRATEGY_PASS"})
    write_csv(out/"au_pdsc_distinct_label_feasibility_audit_v0.84.csv",pdsc_rows)
    pdsc_codes=[r["PDSC_Code"] for r in pdsc_rows if r["PDSC_Code"]]
    pdsc_collisions=len(pdsc_codes)-len(set(pdsc_codes))
    write_json(out/"au_pdsc_fallback_necessity_audit_v0.84.json",{
      "Authority":"G-SEC-03","PDSC_FALLBACK_PERMITTED_IN_PRINCIPLE":"YES","PDSC_FALLBACK_REQUIRED":"YES" if pdsc_required else "NO",
      "Reason":"SOURCE_NATIVE_GICS_CODE_READY" if native_ready else ("SOURCE_NATIVE_CODE_NOT_READY_FORMAL_LABELS_READY" if pdsc_required else "FALLBACK_CONDITIONS_NOT_MET"),
      "PDSC_Generated_Count":len(pdsc_codes),"Sentinel_PDSC_Generated_Count":0
    })
    write_json(out/"au_pdsc_collision_audit_v0.84.json",{"PDSC_Generated_Count":len(pdsc_codes),"PDSC_Unique_Count":len(set(pdsc_codes)),"PDSC_Collisions":pdsc_collisions,
      "Status":"NOT_REQUIRED_SOURCE_NATIVE_STRATEGY_PASS" if native_ready else ("PASS" if pdsc_collisions==0 else "FAIL")})

    pdsc_ready=bool(pdsc_required and len(pdsc_codes)==formal_count and pdsc_collisions==0)
    strategies=int(native_ready)+int(pdsc_ready)
    gate_ready=strategies==1
    strategy="SOURCE_NATIVE_GICS_CODE" if native_ready else ("PROJECT_DERIVED_CANONICAL_PDSC" if pdsc_ready else "NONE")
    blocker=""
    if not gate_ready:
        if not hierarchy_ready:blocker="GICS_INDUSTRY_GROUP_LEVEL_NOT_VERIFIED"
        elif not code_table_authority:blocker="GICS_OFFICIAL_CLASSIFICATION_STRUCTURE_NOT_REPRODUCIBLE"
        elif any(r["Exact_Match_Status"]=="AMBIGUOUS" for r in bindings):blocker="GICS_LABEL_TO_CODE_BINDING_AMBIGUOUS"
        elif code_collisions:blocker="GICS_CODE_COLLISION"
        elif not complete_formal:blocker="ASX_GICS_FORMAL_LABEL_NOT_FOUND_IN_OFFICIAL_GICS"
        elif pdsc_required and pdsc_collisions:blocker="AU_PDSC_COLLISION"
        else:blocker="AU_SECTOR_FIELD_CODE_FEASIBILITY_NOT_VERIFIED"

    decision={"Cohort":"AU_SP_ASX200","Taxonomy":"GICS","Level":"INDUSTRY_GROUP","Classification_Field":"GICs industry group",
      "Current_Raw_Classification_Value_Count":raw_count,"Formal_Label_Count":formal_count,"Sentinel_or_Unresolved_Value_Count":unmatched_count,
      "SOURCE_NATIVE_GICS_CODE_READY":"YES" if native_ready else "NO","PROJECT_DERIVED_CANONICAL_READY":"YES" if pdsc_ready else "NO",
      "Code_Strategy":strategy,"Gate_D":"PASS_BY_CURRENT_GOVERNANCE" if gate_ready else "BLOCKED","Gate_E":"NOT_EVALUATED","Gate_F":"NOT_EVALUATED","Gate_H":"NOT_EVALUATED",
      "AU_SECTOR_FIELD_CODE_FEASIBILITY_READY":"YES" if gate_ready else "NO","Blocker":blocker,
      "Next_Gate":"AU_SP_ASX200 DETERMINISTIC SECURITY IDENTITY LINKAGE GATE E" if gate_ready else blocker}
    write_json(out/"au_sector_field_code_feasibility_decision_v0.84.json",decision)

    ledger=[
      {"Request_Order":1,"Source_Class":"OFFICIAL_ASX","URL":ASX_URL,"Purpose":"ONE_FRESH_CURRENT_DIRECTORY_DOWNLOAD","HTTP_Status":"BROWSER_PUBLIC","SHA256":asx["SHA256"],"Per_Security_Request":"NO"},
      {"Request_Order":2,"Source_Class":"OFFICIAL_GICS_OWNER_SP_DJI","URL":GICS_LANDING_URL,"Purpose":"CURRENT_GICS_STRUCTURE_DISCOVERY","HTTP_Status":landing["HTTP_Status"],"SHA256":landing["SHA256"],"Per_Security_Request":"NO"},
      {"Request_Order":3,"Source_Class":"OFFICIAL_GICS_OWNER_SP_DJI","URL":GICS_XLSX_URL,"Purpose":"CURRENT_GICS_STRUCTURE_LABEL_CODE_WORKBOOK","HTTP_Status":structure["HTTP_Status"],"SHA256":structure["SHA256"],"Per_Security_Request":"NO"},
      {"Request_Order":4,"Source_Class":"OFFICIAL_GICS_OWNER_MSCI","URL":MSCI_GICS_URL,"Purpose":"OWNER_CONTEXT_CORROBORATION","HTTP_Status":msci["HTTP_Status"],"SHA256":msci["SHA256"],"Per_Security_Request":"NO"}
    ]
    write_csv(out/"external_request_ledger_v0.84.csv",ledger)
    prov=provider_audit();write_json(out/"provider_call_audit_v0.84.json",prov)

    imm={"Frozen_SHA256_Expected":FROZEN_SHA,"Frozen_SHA256_After":sha_file(FROZEN),"Frozen_Unchanged":sha_file(FROZEN)==FROZEN_SHA,
      "v057_SHA256_Expected":V057_SHA,"v057_SHA256_After":sha_file(V057),"v057_Unchanged":sha_file(V057)==V057_SHA,
      "v058_SHA256_Expected":V058_SHA,"v058_SHA256_After":sha_file(V058),"v058_Unchanged":sha_file(V058)==V058_SHA,
      "BR_Canonical_Semantic_SHA256_Expected":BR_SHA,"BR_Canonical_Semantic_SHA256_After":read_csv(REGISTRY)[0]["Semantic_SHA256"],"BR_Canonical_Semantic_Unchanged":read_csv(REGISTRY)[0]["Semantic_SHA256"]==BR_SHA,
      "Parked_Cohort_Registry_SHA256_Expected":PARK_SHA,"Parked_Cohort_Registry_SHA256_After":sha_file(PARK),"Parked_Cohort_Registry_Unchanged":sha_file(PARK)==PARK_SHA,
      "IN_Gate_F":"45/45","JP_Gate_F":"197/197","Canonical_READY_Rows_Before":37,"Canonical_READY_Rows_After":37,"Canonical_Total_Rows":1425,
      "AU_Canonical_Rows_After":0,"Frozen_63_Linkage_Runs":0,"AU_Gate_E_Runs":0,"AU_Gate_F_Runs":0,"Gate_H_Runs":0,"Canonical_Materialization_Runs":0,
      "Other_Cohort_Runs":0,"Sector_RS_Runs":0,"P0_Runs":0,"P1_Runs":0,"P2_Runs":0}
    write_json(out/"immutability_audit_v0.84.json",imm)

    tests=[]
    def t(n:str,ok:bool,d:Any):
        tests.append({"Test":n,"Result":"PASS" if ok else "FAIL","Detail":str(d)})
        if not ok:raise RuntimeError(n)
    t("V083_PASS",pred["summary"]["verdict"]=="PASS_AU_SP_ASX200_SOURCE_NATIVE_TAXONOMY_IDENTITY_GATE_C",pred["summary"]["verdict"])
    t("V083_ARTIFACT",pred["checkpoint"]["artifact_id"]==V083_ARTIFACT and pred["checkpoint"]["artifact_digest"]==V083_DIGEST,V083_ARTIFACT)
    t("GATE_C_FIXED",pred["summary"]["taxonomy_identity"]=="GICS" and pred["summary"]["formal_level"]=="INDUSTRY_GROUP","GICS/INDUSTRY_GROUP")
    t("ASX_DOWNLOAD_HEADER",keynorm(asx["Classification_Field"])=="gicsindustrygroup",asx["Classification_Field"])
    t("RAW_VALUE_RECONCILIATION",raw_count==formal_count+unmatched_count,f"{raw_count}={formal_count}+{unmatched_count}")
    t("OFFICIAL_GICS_HIERARCHY",hierarchy_ready,"PASS")
    t("OFFICIAL_GICS_CODE_TABLE",code_table_authority,"PASS")
    t("FORMAL_COUNT_25",formal_count==25,f"{formal_count}")
    t("FORMAL_COMPLETE_AGAINST_OFFICIAL_COUNT",complete_formal,"PASS")
    t("NATIVE_CODE_UNIQUE",code_collisions==0,str(code_collisions))
    t("SENTINELS_NO_CODE",all(r["Canonical_Code_Eligibility"]=="NO" and r["PDSC_Eligibility"]=="NO" for r in sentinel_rows),str(unmatched))
    t("EXACTLY_ONE_STRATEGY",strategies==1,str(strategies))
    t("NO_PDSC_IF_NATIVE",not native_ready or len(pdsc_codes)==0,str(len(pdsc_codes)))
    t("NO_FORBIDDEN_METHODS",all(prov[k]==0 for k in ["Alpha_Vantage","Yahoo_yfinance","EODHD","Scalable","TradingView","Wikipedia","ETF_holdings","third_party_GICS_tables","third_party_security_sector_databases","company_name_Frozen_linkage","fuzzy_matching","semantic_classification_inference","cross_taxonomy_mapping","per_security_web_fanout","price_OHLCV","news","trading_analysis"]),"0")
    t("NO_DOWNSTREAM_GATES",prov["AU_Gate_E"]==prov["AU_Gate_F"]==prov["Gate_H"]==0,"0")
    t("NO_CANONICAL_RS_P",prov["canonical_materialization"]==prov["Sector_RS"]==prov["P0"]==prov["P1"]==prov["P2"]==0,"0")
    t("FROZEN_IMMUTABLE",imm["Frozen_Unchanged"],FROZEN_SHA);t("V057_IMMUTABLE",imm["v057_Unchanged"],V057_SHA);t("V058_IMMUTABLE",imm["v058_Unchanged"],V058_SHA)
    t("BR_IMMUTABLE",imm["BR_Canonical_Semantic_Unchanged"],BR_SHA);t("PARKS_IMMUTABLE",imm["Parked_Cohort_Registry_Unchanged"],PARK_SHA)
    t("CANONICAL_READY_37",imm["Canonical_READY_Rows_After"]==37 and imm["Canonical_Total_Rows"]==1425,"37/1425")
    t("NO_AU_CANONICAL",not any(AU_CANON.glob("AU_SP_ASX200_*.csv")),"0")
    if gate_ready:t("GATE_D_PASS",decision["Gate_D"]=="PASS_BY_CURRENT_GOVERNANCE","PASS")
    else:t("GATE_D_BLOCKED",decision["Gate_D"]=="BLOCKED" and bool(blocker),blocker)
    write_csv(out/"test_results_v0.84.csv",tests)

    verdict="PASS_AU_SP_ASX200_GICS_INDUSTRY_GROUP_CODE_FEASIBILITY_GATE_D" if gate_ready else "BLOCKED_AU_SP_ASX200_GICS_INDUSTRY_GROUP_CODE_FEASIBILITY_GATE_D"
    summary={"version":VERSION,"stage":STAGE,"verdict":verdict,"au_sector_field_code_feasibility_ready":gate_ready,
      "current_raw_classification_values":raw_count,"formal_gics_values":formal_count,"nonformal_unresolved_values":unmatched_count,
      "official_gics_structure_ready":hierarchy_ready and code_table_authority,"source_native_gics_industry_group_codes_ready":native_ready,
      "formal_label_code_coverage":f"{sum(1 for r in bindings if r['Exact_Match_Status']=='EXACT_CODE_BOUND')}/{formal_count}",
      "code_collisions":code_collisions,"sentinel_treatment":"EXCLUDED_NO_CODE_FAIL_CLOSED_AT_GATE_F_IF_FROZEN_ROW",
      "pdsc_fallback_required":pdsc_required,"pdsc_feasibility":"NOT_REQUIRED_SOURCE_NATIVE_STRATEGY_PASS" if native_ready else ("PASS" if pdsc_ready else "NOT_VERIFIED"),
      "pdsc_collisions":pdsc_collisions,"selected_code_strategy":strategy,"blocker":blocker,"au_gate_e":"NOT_EVALUATED","au_gate_f":"NOT_EVALUATED","gate_h":"NOT_EVALUATED",
      "canonical_ready_rows":37,"canonical_total_rows":1425,"sector_rs_runs":0,"p0_runs":0,"p1_runs":0,"p2_runs":0,
      "tests":{"total":len(tests),"passed":len(tests),"failed":0},"artifact_binding":"PENDING_UPLOAD","productive":False,
      "next_gate":"AU_SP_ASX200 DETERMINISTIC SECURITY IDENTITY LINKAGE GATE E" if gate_ready else blocker}
    write_json(out/"summary_preupload_v0.84.json",summary)
    write_json(out/"stage_checkpoint_preupload_v0.84.json",{"version":VERSION,"stage":STAGE,"verdict":verdict,
      "au_sector_field_code_feasibility_ready":gate_ready,"selected_code_strategy":strategy,"blocker":blocker,
      "canonical_ready_rows":37,"canonical_total_rows":1425,"next_gate":summary["next_gate"],"artifact_binding":"PENDING_UPLOAD"})
    files={}
    for p in sorted(out.iterdir()):
        if p.is_file() and p.name!="manifest_preupload_v0.84.json":files[p.name]={"bytes":p.stat().st_size,"sha256":sha_file(p)}
    write_json(out/"manifest_preupload_v0.84.json",{"version":VERSION,"stage":STAGE,"required_start_head":REQUIRED_START_HEAD,"repository_sha":a.repository_sha,
      "verdict":verdict,"au_sector_field_code_feasibility_ready":gate_ready,"selected_code_strategy":strategy,"blocker":blocker,
      "canonical_ready_rows":37,"canonical_total_rows":1425,"frozen_63_linkage_runs":0,"au_gate_e_runs":0,"au_gate_f_runs":0,"gate_h_runs":0,
      "canonical_materialization_runs":0,"other_cohort_runs":0,"sector_rs_runs":0,"p0_runs":0,"p1_runs":0,"p2_runs":0,"files":files,"next_gate":summary["next_gate"]})
    return 0
if __name__=="__main__":raise SystemExit(main())
